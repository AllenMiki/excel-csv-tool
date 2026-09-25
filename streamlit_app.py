# -*- coding: utf-8 -*-
"""
Excel G列 JSON数据转CSV工具
基于 Streamlit 的 Web UI
部署到 Streamlit Cloud 完全免费！
"""

import streamlit as st
import pandas as pd
import json
import base64
import openpyxl
from pathlib import Path

# 页面配置
st.set_page_config(
    page_title="Excel转CSV工具",
    page_icon="📊",
    layout="centered"
)

# 标题
st.markdown("""
# 📊 Excel 转 CSV 工具
### Excel G列 JSON数据转换 - 保留A~D、F列前缀
""")

st.divider()


def convert_excel_to_csv(excel_path):
    """
    读取Excel文件，保留A~D、F列，并解析G列JSON数据
    """
    # 定义需要保留的列
    prefix_cols = {
        1: 'STATION_TYPE',    # A列
        2: 'STATION_NAME',    # B列
        3: 'SN',              # C列
        4: 'RESULT',          # D列
        6: 'CREATE_TIME'      # F列
    }
    
    g_col_index = 7  # G列
    
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    
    result_data_list = []
    row_count = 0
    error_count = 0
    
    for row in range(2, ws.max_row + 1):
        # 获取前缀列数据
        row_prefix = {}
        for col_idx, col_name in prefix_cols.items():
            value = ws.cell(row=row, column=col_idx).value
            row_prefix[col_name] = value
        
        # 获取G列JSON数据
        g_cell_value = ws.cell(row=row, column=g_col_index).value
        
        if g_cell_value and str(g_cell_value).strip():
            try:
                data = json.loads(str(g_cell_value))
                
                if isinstance(data, list):
                    for item in data:
                        if isinstance(item, dict):
                            combined = {**row_prefix, **item}
                            result_data_list.append(combined)
                            row_count += 1
                elif isinstance(data, dict):
                    combined = {**row_prefix, **data}
                    result_data_list.append(combined)
                    row_count += 1
                    
            except json.JSONDecodeError:
                error_count += 1
        else:
            # 空行也保留
            result_data_list.append(row_prefix)
    
    wb.close()
    
    return result_data_list, row_count, error_count


def get_csv_download_link(csv_content, filename="output.csv"):
    """
    生成带BOM的CSV下载链接，使用base64编码确保中文正确
    """
    # UTF-8-BOM: \ufeff 是BOM标记，让Excel正确识别中文编码
    bom = '\ufeff'
    csv_with_bom = bom + csv_content
    
    # 使用base64编码
    b64 = base64.b64encode(csv_with_bom.encode('utf-8')).decode()
    
    href = f'<a href="data:application/octet-stream;base64,{b64}" download="{filename}" style="display:inline-block;padding:0.5rem 1rem;background-color:#ff4b4b;color:white;text-decoration:none;border-radius:0.3rem;font-weight:500;text-align:center;width:100%;">⬇️ 下载 CSV 文件</a>'
    
    return href


# 文件上传
uploaded_file = st.file_uploader(
    "📁 选择 Excel 文件 (.xlsx 或 .xls)",
    type=['xlsx', 'xls'],
    help="拖拽或点击选择文件"
)

if uploaded_file:
    # 显示文件信息
    col1, col2 = st.columns(2)
    with col1:
        st.success(f"✅ 已选择: {uploaded_file.name}")
    with col2:
        size_mb = len(uploaded_file.getvalue()) / (1024 * 1024)
        st.info(f"📦 大小: {size_mb:.2f} MB")
    
    # 转换按钮
    if st.button("🚀 开始转换", type="primary", use_container_width=True):
        with st.spinner("正在转换..."):
            try:
                # 保存上传的文件
                upload_dir = Path("/tmp/temp_uploads")
                upload_dir.mkdir(parents=True, exist_ok=True)
                temp_path = upload_dir / uploaded_file.name
                
                with open(temp_path, "wb") as f:
                    f.write(uploaded_file.getbuffer())
                
                # 执行转换
                result_data, row_count, error_count = convert_excel_to_csv(temp_path)
                
                if result_data:
                    # 转换为DataFrame
                    df = pd.DataFrame(result_data)
                    
                    # 调整列顺序
                    prefix_column_order = ['STATION_TYPE', 'STATION_NAME', 'SN', 'RESULT', 'CREATE_TIME']
                    json_columns = [col for col in df.columns if col not in prefix_column_order]
                    new_column_order = prefix_column_order + json_columns
                    new_column_order = [col for col in new_column_order if col in df.columns]
                    df = df[new_column_order]
                    
                    # 生成CSV - 使用utf-8生成（不带BOM，后面手动添加）
                    csv_content = df.to_csv(index=False, encoding='utf-8')
                    
                    # 生成下载按钮
                    output_filename = f"{Path(uploaded_file.name).stem}_result_data.csv"
                    
                    # 先删除上传的源文件（数据保密）
                    try:
                        if temp_path.exists():
                            temp_path.unlink()
                    except:
                        pass
                    
                    st.success(f"🎉 转换成功！共 {len(df)} 行 × {len(df.columns)} 列")
                    
                    # 统计信息
                    with st.expander("📊 转换统计"):
                        st.write(f"- 有效数据行: {row_count}")
                        st.write(f"- JSON解析失败: {error_count} 行")
                        st.write(f"- 总列数: {len(df.columns)}")
                    
                    # 使用自定义HTML下载链接（确保中文正确）
                    st.markdown(get_csv_download_link(csv_content, output_filename), unsafe_allow_html=True)
                    
                    # 预览数据
                    with st.expander("👁️ 数据预览 (前10行)"):
                        st.dataframe(df.head(10), use_container_width=True)
                        
                else:
                    st.error("❌ 未找到有效数据")
                    
                    # 即使转换失败也删除源文件
                    try:
                        if temp_path.exists():
                            temp_path.unlink()
                    except:
                        pass
                    
            except Exception as e:
                st.error(f"❌ 转换失败: {str(e)}")

# 使用说明
st.divider()
with st.expander("📖 使用说明"):
    st.markdown("""
    ### 使用步骤：
    1. 点击上方「选择文件」按钮，或拖拽Excel文件到上传区域
    2. 点击「开始转换」按钮
    3. 点击「下载CSV文件」按钮保存结果
    
    ### 功能说明：
    - 自动读取Excel文件的G列（RESULT_DATA）
    - 保留A~D列和F列作为前缀字段
    - 将JSON数据展开为平铺的CSV格式
    """)

# 页脚
st.markdown("---")
st.markdown(
    "<div style='text-align: center; color: gray;'>"
    "📊 Excel G列 JSON数据转CSV工具 · 基于 Streamlit 构建"
    "</div>",
    unsafe_allow_html=True
)